"""El .xlsx generado tiene que verse igual al Extractos_cobranzas.xlsx de antes.

Los valores vienen de medir el archivo original: encabezado Calibri 11 negrita
blanco sobre azul 1F4E78 centrado, anchos de columna, fecha en mm-dd-yy e
importe en #,##0.00. Si alguien los cambia sin querer, estos tests avisan.
"""

from __future__ import annotations

import datetime as dt
import io
import sys
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import openpyxl  # noqa: E402

from clasificador.excel import ENCABEZADOS, Movimiento  # noqa: E402
from clasificador.libro import armar_libro  # noqa: E402
from clasificador.reglas import cargar_reglas  # noqa: E402

MOVIMIENTOS = [
    Movimiento(
        fecha=dt.date(2026, 7, 1),
        banco="Galicia",
        importe=150000.0,
        extracto="Transferencia De Terceros",
        detalle="pagare",
        categoria="Cobranza financiamiento clientes (pagaré)",
    ),
    Movimiento(
        fecha=dt.date(2026, 7, 2),
        banco="Roela",
        importe=400000.01,
        extracto="Pared Victor Hugo",
        detalle="TQR",
        categoria="Cobranza cartera planes",
    ),
]


class TestFormatoDelLibro(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        juego = cargar_reglas()
        cls.categorias = juego.categorias_validas
        datos = armar_libro(MOVIMIENTOS, cls.categorias)
        cls.libro = openpyxl.load_workbook(io.BytesIO(datos))
        cls.ws = cls.libro["Cobranzas"]

    def test_hojas(self):
        self.assertEqual(self.libro.sheetnames, ["Cobranzas", "Listas"])

    def test_encabezados_con_acento(self):
        fila = [c.value for c in self.ws[1]]
        self.assertEqual(fila, ENCABEZADOS)

    def test_estilo_del_encabezado(self):
        for celda in self.ws[1]:
            with self.subTest(celda=celda.coordinate):
                self.assertTrue(celda.font.bold)
                self.assertEqual(celda.font.name, "Calibri")
                self.assertEqual(celda.font.size, 11)
                self.assertEqual(celda.font.color.rgb, "FFFFFFFF")
                self.assertEqual(celda.fill.fgColor.rgb, "FF1F4E78")
                self.assertEqual(celda.alignment.horizontal, "center")

    def test_anchos_de_columna(self):
        esperados = {"A": 12.0, "C": 15.0, "D": 60.0, "E": 30.0, "F": 38.0}
        for columna, ancho in esperados.items():
            with self.subTest(columna=columna):
                self.assertEqual(self.ws.column_dimensions[columna].width, ancho)

    def test_formatos_de_numero(self):
        self.assertEqual(self.ws["A2"].number_format, "mm-dd-yy")
        self.assertEqual(self.ws["C2"].number_format, "#,##0.00")

    def test_fecha_e_importe_no_son_texto(self):
        """Si se escriben como texto, Excel no los ordena ni los filtra."""
        self.assertIsInstance(self.ws["A2"].value, (dt.date, dt.datetime))
        self.assertIsInstance(self.ws["C2"].value, (int, float))

    def test_desplegable_sin_igual_adelante(self):
        """Con '=' adelante Excel ofrece reparar el archivo y tira la validacion."""
        validaciones = self.ws.data_validations.dataValidation
        self.assertEqual(len(validaciones), 1)
        dv = validaciones[0]
        self.assertEqual(dv.type, "list")
        self.assertFalse(dv.formula1.startswith("="))
        self.assertEqual(dv.formula1, f"Listas!$A$1:$A${len(self.categorias)}")

    def test_fila_1_congelada_y_autofiltro(self):
        self.assertEqual(self.ws.freeze_panes, "A2")
        self.assertEqual(self.ws.auto_filter.ref, f"A1:F{self.ws.max_row}")

    def test_listas_tiene_las_categorias(self):
        valores = [f[0] for f in self.libro["Listas"].iter_rows(values_only=True)]
        self.assertEqual(valores, list(self.categorias))


if __name__ == "__main__":
    unittest.main()


class TestVolcarEnCarpeta(unittest.TestCase):
    """El camino de OneDrive: escribir el extracto en una carpeta del disco."""

    def setUp(self):
        import tempfile

        self.juego = cargar_reglas()
        self.tmp = tempfile.TemporaryDirectory()
        self.carpeta = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def _volcar(self, movimientos, **kw):
        from clasificador.libro import volcar_en_carpeta

        return volcar_en_carpeta(self.carpeta, movimientos, self.juego, **kw)

    def test_crea_el_archivo(self):
        destino, agregados, previos = self._volcar(MOVIMIENTOS)
        self.assertTrue(destino.exists())
        self.assertEqual(destino.name, "Extractos_cobranzas.xlsx")
        self.assertEqual((agregados, previos), (2, 0))

    def test_es_idempotente(self):
        """Volver a correrlo con el mismo origen no duplica filas."""
        self._volcar(MOVIMIENTOS)
        _, agregados, previos = self._volcar(MOVIMIENTOS)
        self.assertEqual((agregados, previos), (0, 2))

    def test_respeta_las_correcciones_a_mano(self):
        """Lo que el contador corrige en la planilla no se puede pisar."""
        destino, _, _ = self._volcar(MOVIMIENTOS)
        libro = openpyxl.load_workbook(destino)
        libro["Cobranzas"]["F2"] = "Saldo cancelatorio"
        libro.save(destino)

        self._volcar(MOVIMIENTOS)
        ws = openpyxl.load_workbook(destino)["Cobranzas"]
        self.assertEqual(ws["F2"].value, "Saldo cancelatorio")

    def test_no_deja_el_temporal(self):
        """Se escribe a un .tmp y se renombra, para que OneDrive no suba a medias."""
        self._volcar(MOVIMIENTOS)
        self.assertEqual(
            [p.name for p in self.carpeta.iterdir()], ["Extractos_cobranzas.xlsx"]
        )

    def test_carpeta_inexistente_da_error_entendible(self):
        from clasificador.libro import volcar_en_carpeta

        with self.assertRaises(RuntimeError) as ctx:
            volcar_en_carpeta(self.carpeta / "no-existe", MOVIMIENTOS, self.juego)
        self.assertIn("OneDrive", str(ctx.exception))
