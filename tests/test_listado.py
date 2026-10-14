"""Modulo 4: el codigo de Cliente, el Canal y el armado de la hoja H1."""

from __future__ import annotations

import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tests"))

import fixtures  # noqa: E402

from clasificador.cobranza import leer_cobranza  # noqa: E402
from clasificador.config import cargar_config  # noqa: E402
from clasificador.extracto import MovimientoBancario  # noqa: E402
from clasificador.listado import (  # noqa: E402
    anexar_en_listado,
    aprender_codigos,
    armar_listado,
    canal_de,
    codigo_cliente,
    escribir_listado,
    leer_ya_exportados,
    ultima_fila_con_datos,
)


def _historial(ruta: Path) -> Path:
    """Listado banco con el formato real: codigo de 9 digitos, ss y Cuenta vacias."""
    f = dt.datetime
    return fixtures.libro_listado(
        ruta,
        [
            # contrato 216 con prefijo 0 y dos sufijos distintos
            [None, None, "000002161", 5500.0, f(2026, 1, 6), f(2026, 1, 5), f(2026, 1, 6), "PF", None],
            [None, None, "000002162", 5500.0, f(2026, 2, 6), f(2026, 2, 5), f(2026, 2, 6), "PF", None],
            [None, None, "000002161", 5500.0, f(2026, 3, 6), f(2026, 3, 5), f(2026, 3, 6), "RP", None],
            # contrato 1307 con prefijo 1
            [None, None, "100013071", 5000.0, f(2026, 1, 6), f(2026, 1, 5), f(2026, 1, 6), "TR RIO", None],
            [None, None, "100013071", 7000.0, f(2026, 2, 6), f(2026, 2, 5), f(2026, 2, 6), "TR RIO", None],
        ],
    )


class TestCodigoCliente(unittest.TestCase):
    def test_casos_medidos_en_los_archivos_reales(self):
        """Pares contrato/codigo tomados del bloque SIRO del libro de origen."""
        self.assertEqual(codigo_cliente("1313", "1", "1"), "100013131")
        self.assertEqual(codigo_cliente("1012", "0", "5"), "000010125")
        self.assertEqual(codigo_cliente("2118", "0", "1"), "000021181")

    def test_el_tramo_del_medio_es_el_contrato(self):
        """La regla verificada: los 7 digitos del medio son el contrato."""
        codigo = codigo_cliente("1307", "1", "1")
        self.assertEqual(codigo, "100013071")
        self.assertEqual(int(codigo[1:8]), 1307)

    def test_siempre_nueve_caracteres(self):
        for contrato in ("1", "81", "886", "4349"):
            with self.subTest(contrato=contrato):
                self.assertEqual(len(codigo_cliente(contrato, "0", "1")), 9)

    def test_sin_contrato(self):
        self.assertEqual(codigo_cliente("", "0", "1"), "")


class TestAprenderCodigos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = cargar_config()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = _historial(Path(cls.tmp.name) / "listado.xlsx")
        cls.codigos = aprender_codigos(cls.ruta, cls.cfg)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_aprende_el_prefijo_de_cada_contrato(self):
        self.assertEqual(self.codigos.prefijos["216"], "0")
        self.assertEqual(self.codigos.prefijos["1307"], "1")

    def test_el_combo_mas_frecuente_gana(self):
        self.assertEqual(self.codigos.combo_de("216", "1", "1")[:2], ("0", "1"))
        self.assertEqual(self.codigos.combo_de("1307", "0", "9")[:2], ("1", "1"))

    def test_marca_la_procedencia(self):
        self.assertEqual(self.codigos.combo_de("216", "1", "1")[2], "historial")
        self.assertEqual(self.codigos.combo_de("9999", "1", "1")[2], "defecto")

    def test_contrato_desconocido_usa_los_valores_por_defecto(self):
        self.assertEqual(self.codigos.combo_de("9999", "5", "7")[:2], ("5", "7"))

    def test_segundo_pago_toma_otro_sufijo_del_historial(self):
        prefijo, sufijo, procedencia = self.codigos.combo_de("216", "1", "1", orden=2)
        self.assertEqual(prefijo, "0")
        self.assertEqual(sufijo, "2")
        self.assertIn("repetido", procedencia)

    def test_ultima_fila_con_datos(self):
        self.assertEqual(ultima_fila_con_datos(self.ruta, self.cfg), 6)

    def test_ya_exportados(self):
        ya = leer_ya_exportados(self.ruta, self.cfg)
        self.assertIn(("000002161", 5500.0, dt.date(2026, 1, 5)), ya)


class TestCanal(unittest.TestCase):
    def setUp(self):
        self.cfg = cargar_config()

    def test_transferencia_por_banco(self):
        self.assertEqual(canal_de("rio", "", self.cfg), ("TR RIO", ""))
        self.assertEqual(canal_de("galicia", "", self.cfg), ("TR GALICIA", ""))
        self.assertEqual(canal_de("patagonia", "", self.cfg), ("TR PATAGONIA", ""))

    def test_el_codigo_de_la_red_manda(self):
        """Un cobro por SIRO sale con su canal (PF, RP, TQR...), no con el banco."""
        self.assertEqual(canal_de("siro", "PF", self.cfg), ("PF", ""))
        self.assertEqual(canal_de("roela", "TQR", self.cfg), ("TQR", ""))

    def test_siro_sin_codigo_avisa(self):
        canal, aviso = canal_de("siro", "", self.cfg)
        self.assertEqual(canal, "")
        self.assertIn("codigo de canal", aviso)

    def test_medio_desconocido_avisa(self):
        canal, aviso = canal_de("lo que sea", "", self.cfg)
        self.assertEqual(canal, "")
        self.assertIn("no se como traducir", aviso)


class TestArmarListado(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = cargar_config()
        cls.tmp = tempfile.TemporaryDirectory()
        tmp = Path(cls.tmp.name)
        cls.listado = _historial(tmp / "listado.xlsx")
        cls.codigos = aprender_codigos(cls.listado, cls.cfg)
        cls.cobranza = fixtures.libro_cobranza(
            tmp / "cobranza.xlsx",
            {
                "Octubre": [
                    {"nombre": "Perez Juan", "contrato": 216, "cuota": 5500,
                     "fecha_pago": dt.datetime(2026, 10, 5), "monto_pago": 5500,
                     "medio_pago": "rio"},
                    {"nombre": "Parcial Paula", "contrato": 1307, "cuota": 10000,
                     "fecha_pago": dt.datetime(2026, 10, 6), "monto_pago": 3000,
                     "medio_pago": "galicia"},
                    {"nombre": "Impago Ines", "contrato": 300, "cuota": 10000},
                    {"nombre": "Baja Bruno", "contrato": 400, "cuota": 10000,
                     "estado": 3, "fecha_pago": "baja"},
                ]
            },
        )
        cls.base, _ = leer_cobranza(cls.cobranza, cls.cfg)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _armar(self, **kw):
        datos = dict(proceso=dt.date(2026, 10, 14))
        datos.update(kw)
        return armar_listado(self.base, self.cfg, self.codigos, **datos)

    def test_exporta_los_cobros_y_no_los_impagos_ni_las_bajas(self):
        filas = self._armar()
        self.assertEqual([f.contrato for f in filas], ["216", "1307"])

    def test_columnas_con_el_formato_de_h1(self):
        fila = self._armar()[0]
        self.assertIsNone(fila.ss)
        self.assertIsNone(fila.cuenta)
        self.assertEqual(fila.cliente, "000002161")
        self.assertEqual(fila.importe, 5500.0)
        self.assertEqual(fila.fecha_proceso, dt.date(2026, 10, 14))
        self.assertEqual(fila.fecha_pago, dt.date(2026, 10, 5))
        self.assertEqual(fila.canal, "TR RIO")
        self.assertIsNone(fila.subtotal)

    def test_incluye_los_parciales_por_defecto(self):
        """Un pago parcial es plata que entro: el destino la tiene que ver."""
        self.assertEqual(len(self._armar()), 2)
        self.assertEqual(len(self._armar(solo_pagados=True)), 1)

    def test_el_codigo_del_extracto_manda(self):
        """Los cobros SIRO traen el codigo armado por el sistema destino."""
        mov = MovimientoBancario(
            fecha=dt.date(2026, 10, 5), monto=5500.0, banco="Roela",
            titular_origen="Perez Juan", cuit_cvu="", contrato_detectado="216",
            concepto="", codigo_canal="PF", codigo_cliente="000002165",
            fecha_acreditacion=dt.date(2026, 10, 9), mes_origen="Octubre",
            anio_origen=2026,
        )
        fila = self._armar(movimientos=[mov])[0]
        self.assertEqual(fila.cliente, "000002165")
        self.assertEqual(fila.procedencia, "extracto")
        self.assertEqual(fila.canal, "PF")
        self.assertEqual(fila.fecha_acreditacion, dt.date(2026, 10, 9))
        self.assertEqual(fila.revisar, "")

    def test_filtra_por_fecha(self):
        self.assertEqual(len(self._armar(desde=dt.date(2026, 10, 6))), 1)
        self.assertEqual(len(self._armar(hasta=dt.date(2026, 10, 5))), 1)

    def test_no_repite_lo_ya_exportado(self):
        ya = {("000002161", 5500.0, dt.date(2026, 10, 5))}
        filas = self._armar(ya_exportados=ya)
        self.assertEqual([f.contrato for f in filas], ["1307"])

    def test_avisa_si_la_fecha_de_pago_es_posterior_al_proceso(self):
        filas = self._armar(proceso=dt.date(2026, 10, 1))
        self.assertTrue(any("posterior a la fecha de proceso" in f.revisar for f in filas))

    def test_mapa_cliente_manda_sobre_el_historial(self):
        import copy

        cfg = copy.deepcopy(self.cfg)
        cfg.datos["listado"]["mapa_cliente"] = {"216": "999999999"}
        filas = armar_listado(
            self.base, cfg, self.codigos, proceso=dt.date(2026, 10, 14)
        )
        self.assertEqual(filas[0].cliente, "999999999")
        self.assertEqual(filas[0].procedencia, "mapa_cliente")

    def test_escribir_listado(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            ruta = escribir_listado(self._armar(), Path(tmp) / "out.xlsx", self.cfg)
            libro = openpyxl.load_workbook(ruta)
            ws = libro[self.cfg.listado["hoja"]]
            self.assertEqual(
                [c.value for c in ws[1]], self.cfg.listado["encabezados"]
            )
            # El codigo va como texto: si no, 000002161 se vuelve 2161.
            self.assertEqual(ws["C2"].value, "000002161")
            self.assertEqual(ws["C2"].number_format, "@")
            self.assertIn("Revisar", libro.sheetnames)
            self.assertIn("Procedencia", libro.sheetnames)
            libro.close()

    def test_anexar_al_final_de_h1(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            destino = Path(tmp) / "anexado.xlsx"
            desde, hasta = anexar_en_listado(
                self._armar(), self.listado, destino, self.cfg
            )
            self.assertEqual((desde, hasta), (7, 8))
            libro = openpyxl.load_workbook(destino, data_only=True)
            ws = libro["H1"]
            self.assertEqual(ws["C7"].value, "000002161")
            self.assertEqual(ws["D7"].value, 5500.0)
            self.assertEqual(ws["F7"].value, dt.datetime(2026, 10, 5))
            # Y lo que ya estaba sigue estando.
            self.assertEqual(ws["C2"].value, "000002161")
            libro.close()


if __name__ == "__main__":
    unittest.main()
