"""Volcado a SharePoint: link compartido, armado del libro y merge.

Todo lo de aca corre sin red y sin credenciales: Graph se reemplaza por una
sesion falsa que devuelve respuestas armadas a mano.
"""

from __future__ import annotations

import datetime as dt
import io
import sys
import unittest
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from clasificador.excel import (  # noqa: E402
    ENCABEZADOS,
    HOJA_DATOS,
    Movimiento,
    leer_movimientos,
)
from clasificador.libro import HOJA_LISTAS, armar_libro  # noqa: E402
from clasificador.reglas import cargar_reglas  # noqa: E402
from clasificador.sharepoint import (  # noqa: E402
    Carpeta,
    ClienteGraph,
    analizar_link,
    codificar_link,
    combinar,
    decodificar_link,
    listar,
    volcar,
)

URL_CARPETA = (
    "https://achconstrucciones246.sharepoint.com/:f:/s/AdministracionyFinanzas/"
    "IgBzJED7ng-gSJ8pFx9XNs9UAS1olb0dYg0_W14PbLMXyEk?e=AcMGK2"
)


def _mov(dia: int, importe: float, extracto: str, categoria: str = "") -> Movimiento:
    return Movimiento(
        fecha=dt.date(2026, 3, dia),
        banco="Macro",
        importe=importe,
        extracto=extracto,
        detalle="",
        categoria=categoria,
    )


class TestCodificarLink(unittest.TestCase):
    def test_formato_u_bang(self):
        codificado = codificar_link(URL_CARPETA)
        self.assertTrue(codificado.startswith("u!"))
        # base64url sin relleno: ni '=', ni '+', ni '/'
        self.assertNotIn("=", codificado)
        self.assertNotIn("+", codificado)
        self.assertNotIn("/", codificado)

    def test_ida_y_vuelta(self):
        self.assertEqual(decodificar_link(codificar_link(URL_CARPETA)), URL_CARPETA)

    def test_ejemplo_de_la_doc(self):
        """El algoritmo tiene que dar lo mismo que el ejemplo C# de Microsoft."""
        url = "https://onedrive.live.com/redir?resid=1231244193912!12&authkey=!Abcdefg"
        self.assertEqual(
            codificar_link(url),
            "u!aHR0cHM6Ly9vbmVkcml2ZS5saXZlLmNvbS9yZWRpcj9yZXNpZD0xMjMxMjQ0MTkzOTEyITEyJmF1dGhrZXk9IUFiY2RlZmc",
        )

    def test_rechaza_token_sin_prefijo(self):
        with self.assertRaises(ValueError):
            decodificar_link("aHR0cHM6Ly9lamVtcGxv")


class TestAnalizarLink(unittest.TestCase):
    def test_link_de_carpeta_real(self):
        link = analizar_link(URL_CARPETA)
        self.assertEqual(link.hostname, "achconstrucciones246.sharepoint.com")
        self.assertEqual(link.tipo, "carpeta")
        # '/s/' es la forma corta de '/sites/'
        self.assertEqual(link.ruta_sitio, "/sites/AdministracionyFinanzas")

    def test_distingue_archivo_de_carpeta(self):
        url = "https://ejemplo.sharepoint.com/:x:/s/Finanzas/EaBcD?e=1"
        self.assertEqual(analizar_link(url).tipo, "excel")

    def test_salta_el_segmento_de_modo(self):
        url = "https://ejemplo.sharepoint.com/:f:/r/sites/Finanzas/EaBcD"
        self.assertEqual(analizar_link(url).ruta_sitio, "/sites/Finanzas")

    def test_teams(self):
        url = "https://ejemplo.sharepoint.com/:f:/t/Obra/EaBcD"
        self.assertEqual(analizar_link(url).ruta_sitio, "/teams/Obra")

    def test_url_invalida(self):
        with self.assertRaises(ValueError):
            analizar_link("no-es-una-url")


class TestArmarLibro(unittest.TestCase):
    """El formato tiene que ser el mismo que arma ``sheets.preparar``."""

    def setUp(self):
        self.juego = cargar_reglas()
        self.movimientos = [
            _mov(1, 1000.0, "TRANSF 123", "Cobranza cartera planes"),
            _mov(2, 2500.5, "DEPOSITO 456", "Ventas Directas"),
        ]
        self.datos = armar_libro(self.movimientos, self.juego.categorias_validas)
        self.libro = openpyxl.load_workbook(io.BytesIO(self.datos))

    def test_hojas_y_encabezado(self):
        self.assertEqual(self.libro.sheetnames, [HOJA_DATOS, HOJA_LISTAS])
        ws = self.libro[HOJA_DATOS]
        self.assertEqual([c.value for c in ws[1]], ENCABEZADOS)

    def test_hoja_listas_con_las_categorias(self):
        valores = [c[0] for c in self.libro[HOJA_LISTAS].iter_rows(values_only=True)]
        self.assertEqual(valores, list(self.juego.categorias_validas))

    def test_fila_1_congelada(self):
        self.assertEqual(self.libro[HOJA_DATOS].freeze_panes, "A2")

    def test_desplegable_sobre_la_columna_f(self):
        validaciones = self.libro[HOJA_DATOS].data_validations.dataValidation
        self.assertEqual(len(validaciones), 1)
        dv = validaciones[0]
        self.assertEqual(dv.type, "list")
        self.assertEqual(
            dv.formula1, f"={HOJA_LISTAS}!$A$1:$A${len(self.juego.categorias_validas)}"
        )
        self.assertTrue(str(dv.sqref).startswith("F2:F"))
        # showDropDown esta invertido en OOXML: True esconderia la flechita.
        self.assertFalse(dv.showDropDown)

    def test_fecha_e_importe_quedan_como_valores_y_no_como_texto(self):
        """Si la fecha fuera texto, en Excel no se podria ordenar ni filtrar."""
        ws = self.libro[HOJA_DATOS]
        self.assertIsInstance(ws["A2"].value, dt.datetime)
        self.assertEqual(ws["A2"].value.date(), dt.date(2026, 3, 1))
        self.assertEqual(ws["A2"].number_format, "yyyy-mm-dd")
        # Excel guarda 1000.0 y lo devuelve como int: lo que importa es que sea
        # un numero y no una cadena.
        self.assertIsInstance(ws["C2"].value, (int, float))
        self.assertEqual(ws["C2"].number_format, "#,##0.00")

    def test_lo_relee_el_lector_del_proyecto(self):
        """El ciclo bajar -> leer -> mergear -> subir depende de esto."""
        releidos = leer_movimientos(io.BytesIO(self.datos))
        self.assertEqual(len(releidos), 2)
        self.assertEqual([m.clave for m in releidos], [m.clave for m in self.movimientos])
        self.assertEqual(releidos[1].importe, 2500.5)
        self.assertEqual(releidos[1].categoria, "Ventas Directas")

    def test_libro_vacio(self):
        datos = armar_libro([], self.juego.categorias_validas)
        self.assertEqual(leer_movimientos(io.BytesIO(datos)), [])


class TestCombinar(unittest.TestCase):
    def setUp(self):
        self.juego = cargar_reglas()

    def test_no_duplica_lo_que_ya_estaba(self):
        previos = [_mov(1, 1000.0, "TRANSF 123", "Ventas Directas")]
        nuevos = [_mov(1, 1000.0, "TRANSF 123"), _mov(2, 500.0, "TRANSF 999")]
        todos, agregados = combinar(previos, nuevos, self.juego)
        self.assertEqual(agregados, 1)
        self.assertEqual(len(todos), 2)

    def test_conserva_la_categoria_corregida_a_mano(self):
        """Es el punto del merge: regenerar el archivo no debe pisar correcciones."""
        previos = [_mov(1, 1000.0, "TRANSF 123", "Adjudicaciones")]
        nuevos = [_mov(1, 1000.0, "TRANSF 123", "Ventas Directas")]
        todos, agregados = combinar(previos, nuevos, self.juego)
        self.assertEqual(agregados, 0)
        self.assertEqual(todos[0].categoria, "Adjudicaciones")

    def test_clasifica_los_que_vienen_sin_categoria(self):
        nuevos = [_mov(3, 10.0, "X")]
        nuevos[0].detalle = "pagare"
        todos, _ = combinar([], nuevos, self.juego)
        self.assertEqual(todos[0].categoria, self.juego.clasificar("pagare").categoria)

    def test_repetidos_legitimos_no_se_descartan(self):
        """Dos cobros identicos el mismo dia son dos cobros, no un duplicado."""
        nuevos = [_mov(1, 1000.0, "TRANSF 123"), _mov(1, 1000.0, "TRANSF 123")]
        todos, agregados = combinar([], nuevos, self.juego)
        self.assertEqual(agregados, 2)
        self.assertEqual(len(todos), 2)

    def test_reclasificar_todo_pisa_la_categoria_del_origen(self):
        nuevos = [_mov(1, 10.0, "X", "Categoria vieja")]
        nuevos[0].detalle = "pagare"
        todos, _ = combinar([], nuevos, self.juego, respetar_categoria_existente=False)
        self.assertNotEqual(todos[0].categoria, "Categoria vieja")


# --- Graph simulado --------------------------------------------------------
class RespuestaFalsa:
    def __init__(self, status_code: int = 200, cuerpo=None, content: bytes = b""):
        self.status_code = status_code
        self._cuerpo = cuerpo if cuerpo is not None else {}
        self.content = content
        self.text = str(cuerpo)

    def json(self):
        return self._cuerpo


class SesionFalsa:
    """Reemplaza a requests.Session: guarda los pedidos y devuelve lo preparado."""

    def __init__(self, respuestas: dict):
        self.headers: dict[str, str] = {}
        self.respuestas = respuestas
        self.pedidos: list[tuple[str, str, dict]] = []

    def request(self, metodo, url, **kwargs):
        """Las claves son "METODO fragmento-de-la-url"."""
        self.pedidos.append((metodo, url, kwargs))
        for clave, respuesta in self.respuestas.items():
            esperado, _, fragmento = clave.partition(" ")
            if metodo == esperado and fragmento in url:
                return respuesta
        return RespuestaFalsa(404, {"error": {"code": "itemNotFound", "message": url}})


class TestClienteGraph(unittest.TestCase):
    def test_listar_sigue_la_paginacion(self):
        sesion = SesionFalsa(
            {
                "GET /children": RespuestaFalsa(
                    200,
                    {
                        "value": [
                            {
                                "id": "1",
                                "name": "a.xlsx",
                                "size": 10,
                                "lastModifiedDateTime": "2026-01-01T00:00:00Z",
                            },
                            {"id": "sub", "name": "Vieja", "folder": {}},
                        ],
                        "@odata.nextLink": "https://graph.microsoft.com/v1.0/pagina2",
                    },
                ),
                "GET /pagina2": RespuestaFalsa(
                    200,
                    {
                        "value": [
                            {
                                "id": "2",
                                "name": "b.xlsx",
                                "size": 20,
                                "lastModifiedDateTime": "2026-02-01T00:00:00Z",
                            }
                        ]
                    },
                ),
            }
        )
        archivos = listar(ClienteGraph("tok", sesion), Carpeta("d", "i"))
        self.assertEqual([a.nombre for a in archivos], ["a.xlsx", "b.xlsx"])

    def test_el_403_explica_quien_concede_el_permiso(self):
        sesion = SesionFalsa(
            {
                "GET /children": RespuestaFalsa(
                    403, {"error": {"code": "accessDenied", "message": "nope"}}
                )
            }
        )
        with self.assertRaises(RuntimeError) as capturado:
            listar(ClienteGraph("tok", sesion), Carpeta("d", "i"))
        mensaje = str(capturado.exception)
        self.assertIn("Sites.Selected", mensaje)
        self.assertIn("SharePoint Administrator", mensaje)

    def test_el_423_dice_que_cierren_el_excel(self):
        sesion = SesionFalsa(
            {"GET /children": RespuestaFalsa(423, {"error": {"code": "resourceLocked"}})}
        )
        with self.assertRaises(RuntimeError) as capturado:
            listar(ClienteGraph("tok", sesion), Carpeta("d", "i"))
        self.assertIn("abierto en Excel", str(capturado.exception))


class TestVolcar(unittest.TestCase):
    def test_ciclo_completo_sin_red(self):
        """Baja el .xlsx que ya estaba, mergea y sube el nuevo."""
        juego = cargar_reglas()
        previo = armar_libro(
            [_mov(1, 1000.0, "TRANSF 123", "Adjudicaciones")],
            juego.categorias_validas,
        )
        sesion = SesionFalsa(
            {
                "GET /children": RespuestaFalsa(
                    200,
                    {
                        "value": [
                            {
                                "id": "item-1",
                                "name": "Extractos_cobranzas.xlsx",
                                "size": len(previo),
                                "lastModifiedDateTime": "2026-01-01T00:00:00Z",
                                "eTag": '"abc,1"',
                            }
                        ]
                    },
                ),
                "GET /content": RespuestaFalsa(200, {}, content=previo),
                "PUT /content": RespuestaFalsa(200, {"webUrl": "https://sp/x.xlsx"}),
            }
        )
        cliente = ClienteGraph("tok", sesion)
        carpeta = Carpeta("drive-1", "carpeta-1")
        nuevos = [_mov(1, 1000.0, "TRANSF 123"), _mov(5, 777.0, "TRANSF 999")]

        resumen = volcar(cliente, carpeta, nuevos, juego)

        self.assertEqual(resumen.agregados, 1)
        self.assertEqual(resumen.ya_existentes, 1)
        self.assertEqual(resumen.url, "https://sp/x.xlsx")

        put = [p for p in sesion.pedidos if p[0] == "PUT"][0]
        # Reemplaza en lugar de renombrar, y manda el eTag para no pisar
        # ediciones que hayan entrado mientras tanto.
        self.assertEqual(
            put[2]["params"]["@microsoft.graph.conflictBehavior"], "replace"
        )
        self.assertEqual(put[2]["headers"]["if-match"], '"abc,1"')

        subido = leer_movimientos(io.BytesIO(put[2]["data"]))
        self.assertEqual(len(subido), 2)
        self.assertEqual(subido[0].categoria, "Adjudicaciones")

    def test_no_sube_nada_si_no_hay_movimientos_nuevos(self):
        """Subir lo mismo solo ensuciaria el historial de versiones."""
        juego = cargar_reglas()
        previo = armar_libro([_mov(1, 1000.0, "TRANSF 123")], juego.categorias_validas)
        sesion = SesionFalsa(
            {
                "GET /children": RespuestaFalsa(
                    200,
                    {
                        "value": [
                            {
                                "id": "item-1",
                                "name": "Extractos_cobranzas.xlsx",
                                "size": 1,
                                "lastModifiedDateTime": "2026-01-01T00:00:00Z",
                            }
                        ]
                    },
                ),
                "GET /content": RespuestaFalsa(200, {}, content=previo),
            }
        )
        resumen = volcar(
            ClienteGraph("tok", sesion),
            Carpeta("d", "i"),
            [_mov(1, 1000.0, "TRANSF 123")],
            juego,
        )
        self.assertEqual(resumen.agregados, 0)
        self.assertFalse([p for p in sesion.pedidos if p[0] == "PUT"])

    def test_primera_corrida_con_la_carpeta_vacia(self):
        juego = cargar_reglas()
        sesion = SesionFalsa(
            {
                "GET /children": RespuestaFalsa(200, {"value": []}),
                "PUT /content": RespuestaFalsa(201, {"webUrl": "https://sp/x.xlsx"}),
            }
        )
        resumen = volcar(
            ClienteGraph("tok", sesion),
            Carpeta("d", "i"),
            [_mov(1, 1000.0, "TRANSF 123")],
            juego,
        )
        self.assertEqual(resumen.agregados, 1)
        put = [p for p in sesion.pedidos if p[0] == "PUT"][0]
        # Sin archivo previo no hay eTag que mandar.
        self.assertNotIn("if-match", put[2]["headers"])


if __name__ == "__main__":
    unittest.main()
