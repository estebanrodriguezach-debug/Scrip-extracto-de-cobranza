"""Modulo 3: la regla de Estado y el tablero de Mora Activa."""

from __future__ import annotations

import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
sys.path.insert(0, str(RAIZ / "tests"))

import fixtures  # noqa: E402

from clasificador.cobranza import leer_cobranza  # noqa: E402
from clasificador.config import cargar_config  # noqa: E402
from clasificador.estado import (  # noqa: E402
    PAGADO,
    PARCIAL,
    PENDIENTE,
    celdas_de_estado,
    estado_de,
    formula_estado,
    resumen_estados,
)
from clasificador.mora import calcular_mora, escribir_tablero, fin_de_mes  # noqa: E402


class TestEstadoDe(unittest.TestCase):
    def test_los_tres_estados_de_la_especificacion(self):
        self.assertEqual(estado_de(None, 15000), PENDIENTE)
        self.assertEqual(estado_de(0, 15000), PENDIENTE)
        self.assertEqual(estado_de(5000, 15000), PARCIAL)
        self.assertEqual(estado_de(15000, 15000), PAGADO)
        self.assertEqual(estado_de(50000, 15000), PAGADO)

    def test_sin_cuota_de_referencia(self):
        """Si pago y no hay cuota con que comparar, esta pagado."""
        self.assertEqual(estado_de(10000, None), PAGADO)
        self.assertEqual(estado_de(10000, 0), PAGADO)
        self.assertEqual(estado_de(None, None), PENDIENTE)


class TestFormula(unittest.TestCase):
    def setUp(self):
        self.cfg = cargar_config()

    def test_usa_las_columnas_de_la_config(self):
        self.assertEqual(
            formula_estado(3, self.cfg),
            'IF(N(H3)<=0,0,IF(N(H3)>=N(J3),1,"Parcial"))',
        )

    def test_la_formula_dice_lo_mismo_que_la_funcion(self):
        """Si se separan, el tablero diria una cosa y la planilla otra."""
        formula = formula_estado(7, self.cfg)
        for monto, cuota in [(None, 15000), (0, 15000), (5000, 15000),
                             (15000, 15000), (50000, 15000)]:
            with self.subTest(monto=monto, cuota=cuota):
                # Se evalua la formula a mano, con la misma semantica de N().
                m, c = monto or 0, cuota or 0
                excel = 0 if m <= 0 else (1 if m >= c else "Parcial")
                self.assertEqual(excel, estado_de(monto, cuota))
        self.assertIn("N(H7)", formula)


class TestCeldasDeEstado(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = cargar_config()
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = fixtures.libro_cobranza(
            Path(cls.tmp.name) / "c.xlsx",
            {
                "Enero": [
                    {"nombre": "Pago Total", "contrato": 1, "cuota": 10000,
                     "fecha_pago": dt.datetime(2026, 1, 5), "monto_pago": 10000},
                    {"nombre": "Pago Parcial", "contrato": 2, "cuota": 10000,
                     "fecha_pago": dt.datetime(2026, 1, 5), "monto_pago": 3000},
                    {"nombre": "Impago", "contrato": 3, "cuota": 10000},
                    {"nombre": "Baja", "contrato": 4, "cuota": 10000, "estado": 3,
                     "fecha_pago": "baja"},
                    {"nombre": "Nota manual", "contrato": 5, "cuota": 10000,
                     "estado": "pago pero se dio de baja"},
                ]
            },
        )
        cls.base, _ = leer_cobranza(cls.ruta, cls.cfg)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_no_pisa_bajas_ni_notas_manuales(self):
        celdas, conteo = celdas_de_estado(self.base, self.cfg)
        self.assertEqual(conteo["escritas"], 3)
        self.assertEqual(conteo["bajas"], 1)
        self.assertEqual(conteo["manuales"], 1)
        # Las filas 6 y 7 son la baja y la nota manual.
        self.assertEqual(sorted(c.fila for c in celdas), [3, 4, 5])

    def test_escribe_formulas_en_la_columna_estado(self):
        celdas, _ = celdas_de_estado(self.base, self.cfg)
        self.assertTrue(all(c.columna == 12 for c in celdas))
        self.assertTrue(all(c.formula for c in celdas))

    def test_resumen(self):
        reparto = resumen_estados(self.base, self.cfg)
        self.assertEqual(reparto["pagado"], 1)
        self.assertEqual(reparto["parcial"], 1)
        self.assertEqual(reparto["baja"], 1)
        # El impago y la nota manual sin monto cuentan como pendientes.
        self.assertEqual(reparto["pendiente"], 2)


class TestMora(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = cargar_config()
        cls.corte = dt.date(2026, 4, 15)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = fixtures.libro_cobranza(
            Path(cls.tmp.name) / "c.xlsx",
            {
                "Enero": [
                    {"nombre": "Moroso Mario", "contrato": 100, "cuota": 10000,
                     "telefono": 2991111111, "vto_cuota": dt.datetime(2026, 1, 10)},
                    {"nombre": "Parcial Paula", "contrato": 200, "cuota": 10000,
                     "monto_pago": 4000, "fecha_pago": dt.datetime(2026, 1, 5),
                     "vto_cuota": dt.datetime(2026, 1, 10)},
                    {"nombre": "Al Dia Ana", "contrato": 300, "cuota": 10000,
                     "monto_pago": 10000, "fecha_pago": dt.datetime(2026, 1, 5)},
                    {"nombre": "Baja Bruno", "contrato": 400, "cuota": 10000,
                     "estado": 3, "fecha_pago": "baja"},
                ],
                "Febrero": [
                    {"nombre": "Moroso Mario", "contrato": 100, "cuota": 10000,
                     "telefono": 2991111111, "vto_cuota": dt.datetime(2026, 2, 10)},
                ],
                "Mayo": [
                    # Vence despues del corte: todavia no es mora.
                    {"nombre": "Futuro Fabian", "contrato": 500, "cuota": 10000,
                     "vto_cuota": dt.datetime(2026, 5, 10)},
                ],
            },
        )
        base, _ = leer_cobranza(cls.ruta, cls.cfg)
        cls.clientes = calcular_mora(base, cls.cfg, 2026, cls.corte)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def _cliente(self, contrato):
        return next((c for c in self.clientes if c.contrato == contrato), None)

    def test_consolida_cuotas_de_varios_meses(self):
        """Es el sentido del tablero: las cuotas viven en pestañas distintas."""
        mario = self._cliente("100")
        self.assertIsNotNone(mario)
        self.assertEqual(mario.cuotas, 2)
        self.assertEqual(mario.total_vencido, 20000.0)
        self.assertEqual(mario.meses, ["Enero", "Febrero"])
        self.assertEqual(mario.vto_mas_antiguo, dt.date(2026, 1, 10))

    def test_incluye_los_parciales_por_el_saldo(self):
        paula = self._cliente("200")
        self.assertEqual(paula.total_vencido, 6000.0)
        self.assertEqual(paula.pagado_a_cuenta, 4000.0)

    def test_excluye_al_dia_bajas_y_no_vencidos(self):
        self.assertIsNone(self._cliente("300"))
        self.assertIsNone(self._cliente("400"))
        self.assertIsNone(self._cliente("500"))

    def test_dias_de_mora(self):
        self.assertEqual(self._cliente("100").dias_de_mora(self.corte), 95)

    def test_ordenado_por_monto(self):
        montos = [c.total_vencido for c in self.clientes]
        self.assertEqual(montos, sorted(montos, reverse=True))

    def test_telefono_para_poder_llamar(self):
        self.assertEqual(self._cliente("100").telefono, "2991111111")

    def test_fin_de_mes(self):
        self.assertEqual(fin_de_mes("Febrero", 2026), dt.date(2026, 2, 28))
        self.assertEqual(fin_de_mes("Marzo", 2026), dt.date(2026, 3, 31))
        self.assertIsNone(fin_de_mes("Pagare Enero", 2026))

    def test_sin_vto_usa_el_fin_del_mes_de_la_hoja(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = fixtures.libro_cobranza(
                Path(tmp) / "x.xlsx",
                {"Enero": [{"nombre": "Sin Vto", "contrato": 9, "cuota": 5000}]},
            )
            base, _ = leer_cobranza(ruta, self.cfg)
            clientes = calcular_mora(base, self.cfg, 2026, dt.date(2026, 3, 1))
        self.assertEqual(len(clientes), 1)
        self.assertEqual(clientes[0].vto_mas_antiguo, dt.date(2026, 1, 31))

    def test_tablero_se_escribe(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = escribir_tablero(self.clientes, Path(tmp) / "m.xlsx", self.corte)
            import openpyxl

            libro = openpyxl.load_workbook(ruta)
            ws = libro["Mora Activa"]
            self.assertIn("Mora activa al 15/04/2026", str(ws["A1"].value))
            # Las cinco columnas que pide la especificacion, en orden.
            encabezados = [c.value for c in ws[4]][:5]
            self.assertEqual(
                encabezados,
                ["Nombre", "Telefono", "Contrato", "Cuotas Adeudadas",
                 "Monto Total Vencido"],
            )
            libro.close()


if __name__ == "__main__":
    unittest.main()
