"""Libros .xlsx sinteticos para los tests.

Los archivos reales tienen datos de clientes (nombres, telefonos, CUIT) y no se
versionan, asi que los tests arman en memoria libros con la misma forma: el
encabezado de Cobranza General en la fila 2, las columnas A/B/C sin titulo, el
bloque de estadisticas abajo de la tabla, las formulas de Estado y Cantidad
cuotas, y la hoja H1 del Listado banco.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl

ENCABEZADO_COBRANZA = {
    "D": "Telefono", "E": "Localidad", "F": "Provincia", "G": "Fecha de pago",
    "H": "Monto de Pago", "I": "Medio de Pago", "J": "Cuota actual",
    "K": "Mes Inicio", "L": "Estado", "M": "Supervisor", "N": "Vendedor",
    "O": "Vto Cuota", "P": "Cantidad cuotas", "Q": "Obse",
}


def hoja_de_cobranza(ws, filas: list[dict]) -> None:
    """Escribe una hoja mensual con el layout real."""
    for letra, titulo in ENCABEZADO_COBRANZA.items():
        ws[f"{letra}2"] = titulo
    for i, datos in enumerate(filas, start=3):
        ws[f"A{i}"] = datos.get("marca", "A")
        ws[f"B{i}"] = datos.get("nombre", "")
        ws[f"C{i}"] = datos.get("contrato")
        ws[f"D{i}"] = datos.get("telefono", 2990000000)
        ws[f"E{i}"] = datos.get("localidad", "Zapala")
        ws[f"F{i}"] = datos.get("provincia", "Neuquen")
        if datos.get("fecha_pago") is not None:
            ws[f"G{i}"] = datos["fecha_pago"]
        if datos.get("monto_pago") is not None:
            ws[f"H{i}"] = datos["monto_pago"]
        if datos.get("medio_pago"):
            ws[f"I{i}"] = datos["medio_pago"]
        ws[f"J{i}"] = datos.get("cuota")
        ws[f"L{i}"] = datos.get("estado", f"=COUNT(G{i})")
        ws[f"M{i}"] = datos.get("supervisor", "Moreira David")
        ws[f"N{i}"] = datos.get("vendedor", "Lezcano Juan")
        if datos.get("vto_cuota") is not None:
            ws[f"O{i}"] = datos["vto_cuota"]
        ws[f"P{i}"] = f"=COUNTIF(B:B,B{i})"
        if datos.get("obse"):
            ws[f"Q{i}"] = datos["obse"]

    # El bloque de estadisticas que vive abajo de la tabla en el archivo real:
    # mismas columnas, sin nombre de cliente. No tiene que leerse como cliente.
    fin = len(filas) + 5
    ws[f"A{fin}"] = "Cobrado"
    ws[f"C{fin}"] = 63811500
    ws[f"H{fin}"] = 450
    ws[f"I{fin}"] = 82.6466
    ws[f"A{fin + 1}"] = "Pendiente"
    ws[f"C{fin + 1}"] = 12283500


def libro_cobranza(ruta: str | Path, meses: dict[str, list[dict]]) -> Path:
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for mes, filas in meses.items():
        hoja_de_cobranza(libro.create_sheet(mes), filas)
    ruta = Path(ruta)
    libro.save(ruta)
    libro.close()
    return ruta


def libro_ingresos(ruta: str | Path, hojas: dict[str, list[dict]]) -> Path:
    """Libro de Ingresos Bancarios con el layout generico (Rio/Galicia/Macro).

    A extracto, B fecha, C importe, D contrato, E nombre, F detalle.
    """
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for titulo, movimientos in hojas.items():
        ws = libro.create_sheet(titulo)
        for i, mov in enumerate(movimientos, start=1):
            ws[f"A{i}"] = mov.get("concepto", "Transferencia Recibida")
            ws[f"B{i}"] = mov["fecha"]
            ws[f"C{i}"] = mov["monto"]
            ws[f"D{i}"] = mov.get("contrato")
            ws[f"E{i}"] = mov.get("nombre", "")
            ws[f"F{i}"] = mov.get("detalle", "")
    ruta = Path(ruta)
    libro.save(ruta)
    libro.close()
    return ruta


def libro_ingresos_siro(ruta: str | Path, titulo: str, cupones: list[dict]) -> Path:
    """Hoja de Roela con el bloque SIRO a la izquierda.

    A fecha de carga, B importe, C codigo de cliente, D fecha de pago,
    E fecha de acreditacion, F codigo de canal, G contrato, H nombre.
    """
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.title = titulo
    ws["A1"] = "CUPON PAGO FACIL-RAPIPAGO"
    for i, cupon in enumerate(cupones, start=2):
        ws[f"A{i}"] = cupon.get("carga", cupon["pago"])
        ws[f"B{i}"] = cupon["monto"]
        ws[f"C{i}"] = cupon.get("codigo")
        ws[f"D{i}"] = cupon["pago"]
        ws[f"E{i}"] = cupon.get("acreditacion", cupon["pago"])
        ws[f"F{i}"] = cupon.get("canal", "PF")
        ws[f"G{i}"] = cupon.get("contrato")
        ws[f"H{i}"] = cupon.get("nombre", "")
    ruta = Path(ruta)
    libro.save(ruta)
    libro.close()
    return ruta


def libro_listado(ruta: str | Path, filas: list[list]) -> Path:
    """Listado banco con la hoja H1 y sus nueve columnas."""
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.title = "H1"
    ws.append(
        ["ss", "Cuenta", "Cliente", "Importe", "Fecha de Proceso",
         "Fecha de Pago", "Fecha de Acred.", "Canal", "SubTotal"]
    )
    for fila in filas:
        ws.append(fila)
    for f in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for i in (4, 5, 6):
            f[i].number_format = "dd/mm/yyyy"
    ruta = Path(ruta)
    libro.save(ruta)
    libro.close()
    return ruta
