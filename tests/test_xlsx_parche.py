"""Escritura quirurgica: que escriba bien y que no se lleve nada puesto."""

from __future__ import annotations

import datetime as dt
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

import openpyxl  # noqa: E402

from clasificador.xlsx_parche import (  # noqa: E402
    Celda,
    _es_formato_de_fecha,
    _estilos_de_fecha,
    parchear,
)


def _libro_de_prueba(ruta: Path) -> Path:
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.title = "Enero"
    ws.append(["Nombre", "Contrato", "Fecha", "Monto", "Medio", "Estado"])
    for i in range(2, 6):
        ws[f"A{i}"] = f"Cliente {i}"
        ws[f"B{i}"] = 2000 + i
        ws[f"C{i}"] = dt.datetime(2026, 1, i)
        ws[f"C{i}"].number_format = "dd/mm/yyyy"
        ws[f"F{i}"] = f"=COUNT(C{i})"
    otra = libro.create_sheet("Febrero")
    otra.append(["Nombre"])
    otra["A2"] = "Cliente de febrero"
    libro.save(ruta)
    libro.close()
    return ruta


class TestParchear(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.origen = _libro_de_prueba(Path(self.tmp.name) / "base.xlsx")
        self.destino = Path(self.tmp.name) / "salida.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def _leer(self, ruta=None, **kw):
        libro = openpyxl.load_workbook(ruta or self.destino, **kw)
        self.addCleanup(libro.close)
        return libro

    def test_escribe_los_tres_tipos_de_dato(self):
        parchear(
            self.origen,
            self.destino,
            [
                Celda("Enero", 2, 3, dt.date(2026, 3, 15)),
                Celda("Enero", 2, 4, 50000.0),
                Celda("Enero", 2, 5, "rio"),
            ],
        )
        ws = self._leer()["Enero"]
        self.assertEqual(ws["C2"].value, dt.datetime(2026, 3, 15))
        self.assertEqual(ws["D2"].value, 50000.0)
        self.assertEqual(ws["E2"].value, "rio")

    def test_una_fecha_queda_con_formato_de_fecha(self):
        """Sin esto la fecha se ve como 46096 en vez de 15/03/2026.

        Paso de verdad al anexar filas al Listado banco: sus filas vacias ya
        venian preformateadas, pero con formato general.
        """
        parchear(
            self.origen, self.destino, [Celda("Enero", 3, 3, dt.date(2026, 3, 15))]
        )
        celda = self._leer()["Enero"]["C3"]
        self.assertEqual(celda.value, dt.datetime(2026, 3, 15))
        self.assertIn("yy", celda.number_format)

    def test_escribe_una_formula_y_borra_el_valor_viejo(self):
        parchear(
            self.origen,
            self.destino,
            [Celda("Enero", 2, 6, formula='IF(N(D2)<=0,0,1)')],
        )
        ws = self._leer()["Enero"]
        self.assertEqual(ws["F2"].value, "=IF(N(D2)<=0,0,1)")

    def test_no_toca_lo_que_no_se_pidio(self):
        parchear(self.origen, self.destino, [Celda("Enero", 2, 4, 1.0)])
        libro = self._leer()
        self.assertEqual(libro["Enero"]["F3"].value, "=COUNT(C3)")
        self.assertEqual(libro["Enero"]["A3"].value, "Cliente 3")
        self.assertEqual(libro["Febrero"]["A2"].value, "Cliente de febrero")

    def test_conserva_todas_las_partes_del_zip(self):
        """Lo que justifica que este modulo exista."""
        parchear(self.origen, self.destino, [Celda("Enero", 2, 4, 1.0)])
        antes = set(zipfile.ZipFile(self.origen).namelist())
        despues = set(zipfile.ZipFile(self.destino).namelist())
        self.assertEqual(antes - despues, set())

    def test_solo_cambia_la_hoja_tocada(self):
        parchear(self.origen, self.destino, [Celda("Enero", 2, 4, 1.0)])
        a, b = zipfile.ZipFile(self.origen), zipfile.ZipFile(self.destino)
        distintas = {
            n for n in set(a.namelist()) & set(b.namelist()) if a.read(n) != b.read(n)
        }
        # workbook.xml cambia por fullCalcOnLoad; sheet2 (Febrero) no.
        self.assertNotIn("xl/worksheets/sheet2.xml", distintas)

    def test_escribe_en_sitio_sin_perder_el_archivo(self):
        parchear(self.origen, self.origen, [Celda("Enero", 2, 5, "galicia")])
        ws = self._leer(self.origen)["Enero"]
        self.assertEqual(ws["E2"].value, "galicia")
        self.assertFalse(
            (self.origen.parent / f"{self.origen.name}.tmp").exists(),
            "quedo el temporal sin renombrar",
        )

    def test_crea_la_celda_si_no_existia(self):
        # La fila 2 no tiene celda en la columna E ni en la G.
        parchear(
            self.origen,
            self.destino,
            [Celda("Enero", 2, 7, "nota al final"), Celda("Enero", 2, 5, "rio")],
        )
        ws = self._leer()["Enero"]
        self.assertEqual(ws["E2"].value, "rio")
        self.assertEqual(ws["G2"].value, "nota al final")

    def test_texto_con_caracteres_de_xml(self):
        parchear(self.origen, self.destino, [Celda("Enero", 2, 5, 'a & b < c "d"')])
        self.assertEqual(self._leer()["Enero"]["E2"].value, 'a & b < c "d"')

    def test_hoja_inexistente_avisa(self):
        with self.assertRaises(ValueError) as e:
            parchear(self.origen, self.destino, [Celda("Marzo", 2, 3, 1)])
        self.assertIn("Marzo", str(e.exception))

    def test_crea_la_fila_si_no_existia(self):
        """Hace falta para anexar al final de una hoja."""
        parchear(
            self.origen,
            self.destino,
            [Celda("Enero", 9, 1, "nuevo"), Celda("Enero", 8, 3, dt.date(2026, 5, 1))],
        )
        ws = self._leer()["Enero"]
        self.assertEqual(ws["A9"].value, "nuevo")
        self.assertEqual(ws["C8"].value, dt.datetime(2026, 5, 1))
        self.assertEqual(ws["A5"].value, "Cliente 5")  # lo anterior sigue

    def test_filas_creadas_quedan_en_orden(self):
        """Excel repara el archivo si las filas no van ordenadas."""
        parchear(
            self.origen,
            self.destino,
            [Celda("Enero", 12, 1, "doce"), Celda("Enero", 10, 1, "diez")],
        )
        import re
        import zipfile

        with zipfile.ZipFile(self.destino) as zf:
            xml = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        numeros = [int(n) for n in re.findall(r'<row[^>]*\sr="(\d+)"', xml)]
        self.assertEqual(numeros, sorted(numeros))
        self.assertIn(10, numeros)
        self.assertIn(12, numeros)

    def test_sin_celdas_solo_copia(self):
        parchear(self.origen, self.destino, [])
        self.assertTrue(self.destino.exists())
        self.assertEqual(self.origen.read_bytes(), self.destino.read_bytes())


class TestEstilosDeFecha(unittest.TestCase):
    def test_codigos_de_formato(self):
        self.assertTrue(_es_formato_de_fecha("dd/mm/yyyy"))
        self.assertTrue(_es_formato_de_fecha("mm-dd-yy"))
        self.assertFalse(_es_formato_de_fecha("#,##0.00"))
        self.assertFalse(_es_formato_de_fecha("General"))

    def test_no_se_confunde_con_un_literal(self):
        """Una moneda con la letra 'd' adentro de un literal no es fecha."""
        self.assertFalse(_es_formato_de_fecha('"usd"#,##0.00'))
        self.assertFalse(_es_formato_de_fecha('[$-409]#,##0'))

    def test_detecta_los_estilos_del_libro(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = _libro_de_prueba(Path(tmp) / "x.xlsx")
            with zipfile.ZipFile(ruta) as zf:
                estilos = _estilos_de_fecha(zf.read("xl/styles.xml").decode("utf-8"))
            self.assertTrue(estilos, "no encontro el estilo dd/mm/yyyy")


if __name__ == "__main__":
    unittest.main()
