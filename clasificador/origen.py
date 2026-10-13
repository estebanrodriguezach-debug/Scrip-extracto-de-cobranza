"""Lectura del libro "Ingresos Bancarios Clientes", el origen real.

El archivo tiene ~105 hojas, una por banco y mes ("Rio Enero", "Macro Feb 26",
"Galicia oct"...), y el layout NO es uniforme:

* Rio / Galicia / Macro / Frances:  A=extracto, B=fecha, C=importe,
  D=contrato, E=nombre, F=detalle.
* Patagonia: las columnas se corren entre hojas (en "Patagonia Agosto 26" la
  fecha esta en D, en "Patagonia Octubre 26" en C) y el extracto ocupa dos
  columnas que hay que unir.
* Roela: dos tablas una al lado de la otra. La izquierda son cheques
  (A=carga, B=importe, C=nro cheque, D=fecha del cheque, E=acreditacion,
  F=tipo, G=contrato, H=nombre) y la derecha transferencias (L=codigo,
  M=extracto, N=fecha, O=importe, P=contrato, Q=nombre, R=detalle).

Por eso las columnas se detectan por hoja en vez de fijarlas: lo unico
constante en todo el libro es que importe, contrato, nombre y detalle van
pegados a la derecha de la fecha. El bloque de cheques de Roela es la
excepcion y esta declarada aparte.
"""

from __future__ import annotations

import datetime as dt
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from clasificador.excel import Movimiento

BANCOS = ["Patagonia", "Patag", "Pata", "Galicia", "Roela", "Macro", "Rio", "Frances"]
ALIAS_BANCO = {"Patag": "Patagonia", "Pata": "Patagonia"}

# Bloque de cheques de Roela: no sigue la regla de adyacencia. Ademas la fecha
# que vale es la D (fecha del cheque), no la A en que se carga: contrastado
# contra el extracto ya clasificado, D acierta 92% y A solo 14%.
ROELA_CHEQUES = {"fecha": 3, "importe": 1, "detalle": 5, "contrato": 6, "nombre": 7}

FILAS_MUESTRA = 400

# Un año tipeado mal pasa desapercibido: la fila entra igual, con una fecha
# absurda, y el extracto queda ordenado de forma rara. Pasado este margen se
# avisa, pero la fila se conserva: no es el script quien decide tirar datos.
ANIO_MINIMO = 2020
ANIOS_DE_MARGEN = 2


def _sin_acentos(texto: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c)
    )


def banco_de_hoja(titulo: str) -> str | None:
    """Saca el banco del nombre de la hoja ('Macro Feb 26' -> 'Macro')."""
    limpio = _sin_acentos(titulo).strip().lower()
    for banco in BANCOS:
        if limpio.startswith(banco.lower()):
            return ALIAS_BANCO.get(banco, banco)
    return None


@dataclass(frozen=True)
class Bloque:
    """Posiciones (0-based) de cada dato dentro de una hoja."""

    fecha: int
    importe: int
    contrato: int | None = None
    nombre: int | None = None
    detalle: int | None = None
    extracto: tuple[int, ...] = ()


def _es_fecha(valor: object) -> bool:
    return isinstance(valor, dt.datetime)


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def detectar_bloques(filas: list[tuple], ancho: int, banco: str) -> list[Bloque]:
    """Encuentra, por hoja, donde esta cada dato.

    Busca las columnas con fechas y arma un bloque por cada una, asumiendo
    importe/contrato/nombre/detalle a la derecha. Descarta las columnas de
    fecha secundarias (las de Roela, que repiten la fecha del cheque).
    """
    conteo_fechas = [0] * ancho
    conteo_texto = [0] * ancho
    for fila in filas:
        for i, valor in enumerate(fila[:ancho]):
            if _es_fecha(valor):
                conteo_fechas[i] += 1
            elif isinstance(valor, str) and len(valor.strip()) > 12:
                conteo_texto[i] += 1

    if not any(conteo_fechas):
        return []

    minimo = max(2, max(conteo_fechas) * 0.15)
    candidatas = [i for i, n in enumerate(conteo_fechas) if n >= minimo]

    # Una columna de fecha solo abre un bloque si a su derecha hay importes.
    inicios = []
    for i in candidatas:
        if i + 1 >= ancho:
            continue
        numericos = sum(
            1
            for fila in filas
            if isinstance(fila[i + 1] if i + 1 < len(fila) else None, (int, float))
        )
        if numericos >= minimo:
            inicios.append(i)

    bloques = []
    usadas: set[int] = set()
    for i in inicios:
        if any(abs(i - u) <= 2 for u in usadas):
            continue  # fecha repetida del mismo bloque (caso Roela D/E)
        usadas.add(i)
        extracto = tuple(
            c
            for c in range(0, i)
            if conteo_texto[c] >= minimo and not any(abs(c - u) <= 0 for u in usadas)
        )
        bloques.append(
            Bloque(
                fecha=i,
                importe=i + 1,
                contrato=i + 2 if i + 2 < ancho else None,
                nombre=i + 3 if i + 3 < ancho else None,
                detalle=i + 4 if i + 4 < ancho else None,
                extracto=extracto,
            )
        )

    if banco == "Roela":
        bloques = [b for b in bloques if b.fecha > ROELA_CHEQUES["fecha"] + 1]
        bloques.insert(
            0,
            Bloque(
                fecha=ROELA_CHEQUES["fecha"],
                importe=ROELA_CHEQUES["importe"],
                contrato=ROELA_CHEQUES["contrato"],
                nombre=ROELA_CHEQUES["nombre"],
                detalle=ROELA_CHEQUES["detalle"],
                extracto=(),
            ),
        )
    return bloques


def _celda(fila: tuple, i: int | None):
    if i is None or i >= len(fila):
        return None
    return fila[i]


def leer_libro(ruta: str | Path) -> tuple[list[Movimiento], list[str]]:
    """Devuelve (movimientos, avisos) leyendo todas las hojas del libro."""
    libro = openpyxl.load_workbook(ruta, data_only=True)
    movimientos: list[Movimiento] = []
    avisos: list[str] = []
    try:
        for ws in libro.worksheets:
            banco = banco_de_hoja(ws.title)
            if banco is None:
                avisos.append(f"Hoja {ws.title!r}: no reconozco el banco, la salteo.")
                continue

            filas = [
                f
                for f in ws.iter_rows(values_only=True)
                if any(v is not None for v in f)
            ]
            if not filas:
                continue
            ancho = max(len(f) for f in filas)
            bloques = detectar_bloques(filas[:FILAS_MUESTRA], ancho, banco)
            if not bloques:
                avisos.append(
                    f"Hoja {ws.title!r}: no encontre columna de fecha, la salteo."
                )
                continue

            for fila in filas:
                for bloque in bloques:
                    fecha = _celda(fila, bloque.fecha)
                    importe = _celda(fila, bloque.importe)
                    if not _es_fecha(fecha):
                        continue
                    if not isinstance(importe, (int, float)):
                        continue

                    partes = [_texto(_celda(fila, c)) for c in bloque.extracto]
                    extracto = " ".join(p for p in partes if p and not p.isdigit())
                    nombre = _texto(_celda(fila, bloque.nombre))
                    if not extracto:
                        # Los cheques de Roela no traen texto del banco: el
                        # extracto original usa el nombre del cliente.
                        extracto = nombre

                    if not (
                        ANIO_MINIMO <= fecha.year <= dt.date.today().year + ANIOS_DE_MARGEN
                    ):
                        avisos.append(
                            f"Hoja {ws.title!r}: la fecha {fecha.date()} no parece "
                            "real, revisa esa fila en el archivo de origen. "
                            "La dejo igual, sin corregir."
                        )

                    movimientos.append(
                        Movimiento(
                            fecha=fecha.date(),
                            banco=banco,
                            importe=float(importe),
                            extracto=extracto,
                            detalle=_texto(_celda(fila, bloque.detalle)),
                        )
                    )
        return movimientos, avisos
    finally:
        libro.close()


def detectar_formato(ruta: str | Path) -> str:
    """'cobranzas' si el libro ya tiene el formato del extracto, si no 'ingresos'."""
    libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
    try:
        if "Cobranzas" in libro.sheetnames:
            ws = libro["Cobranzas"]
            primera = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            cabeceras = {str(v).strip().lower() for v in primera if v}
            if {"fecha", "banco", "importe", "extracto"} <= cabeceras:
                return "cobranzas"
        return "ingresos"
    finally:
        libro.close()


def leer(ruta: str | Path, formato: str = "auto") -> tuple[list[Movimiento], list[str]]:
    """Lee cualquiera de los dos formatos y devuelve (movimientos, avisos)."""
    from clasificador.excel import leer_movimientos

    if formato == "auto":
        formato = detectar_formato(ruta)
    if formato == "cobranzas":
        return leer_movimientos(ruta), []
    return leer_libro(ruta)
