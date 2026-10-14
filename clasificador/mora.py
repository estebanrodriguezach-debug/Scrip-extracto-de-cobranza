"""MODULO 3 (parte 2) - Tablero de Mora Activa.

Consolida, por cliente, todo lo que quedo impago o pago a medias y ya vencido.
Es la lista con la que el equipo sale a llamar, asi que trae telefono y, de
yapa, supervisor y vendedor para saber a quien preguntarle.

Dos decisiones que cambian el resultado y conviene tener presentes:

* **Se cruza todo el año, no un mes.** "Cuotas Adeudadas" no tiene sentido
  dentro de una sola hoja: el cliente que no pago marzo ni abril tiene dos
  cuotas, y aparecen en dos pestañas distintas. Se recorren todas las hojas
  mensuales y se agrupa por contrato.
* **Si falta ``Vto Cuota``, se usa el fin del mes de la hoja.** En el archivo
  real esa columna esta cargada en apenas un 40% de las filas. Tomarla como
  unico criterio dejaria afuera a la mayoria de los morosos; usar el mes de la
  pestaña es lo que el equipo asume cuando mira la planilla. Se puede apagar
  con ``usar_fin_de_mes_si_falta_vto`` en ``conciliacion.json``.

El tablero se escribe como **archivo aparte**, no como una pestaña nueva en
``Cobranza General``. Agregar una hoja a ese libro obligaria a tocar su
estructura interna (``workbook.xml``, los rels, ``[Content_Types].xml``), y es
justo lo que no conviene hacerle a un archivo con 1.119 comentarios adentro.
Como el tablero se regenera entero en cada corrida, tenerlo suelto ademas es
mas practico: se comparte, se filtra y se imprime sin tocar la base.
"""

from __future__ import annotations

import calendar
import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from clasificador.cobranza import FilaCobranza
from clasificador.config import Config
from clasificador.estado import PARCIAL, PENDIENTE, estado_de
from clasificador.extracto import MESES_DEL_ANIO

ENCABEZADOS_MORA = [
    # Las cinco que pide la especificacion, primero y en orden.
    "Nombre",
    "Telefono",
    "Contrato",
    "Cuotas Adeudadas",
    "Monto Total Vencido",
    # Lo que hace falta para poder trabajar la lista.
    "Pagado a cuenta",
    "Meses adeudados",
    "Vto mas antiguo",
    "Dias de mora",
    "Localidad",
    "Provincia",
    "Supervisor",
    "Vendedor",
]


def fin_de_mes(mes: str, anio: int) -> dt.date | None:
    """Ultimo dia del mes que nombra la hoja ('Marzo' -> 31/03)."""
    if mes not in MESES_DEL_ANIO:
        return None
    numero = MESES_DEL_ANIO.index(mes) + 1
    return dt.date(anio, numero, calendar.monthrange(anio, numero)[1])


@dataclass
class ClienteEnMora:
    """Un cliente con todo lo que debe, sumado a traves de los meses."""

    contrato: str
    nombre: str
    telefono: str = ""
    localidad: str = ""
    provincia: str = ""
    supervisor: str = ""
    vendedor: str = ""
    cuotas: int = 0
    total_vencido: float = 0.0
    pagado_a_cuenta: float = 0.0
    meses: list[str] = field(default_factory=list)
    vto_mas_antiguo: dt.date | None = None

    def dias_de_mora(self, corte: dt.date) -> int:
        if self.vto_mas_antiguo is None:
            return 0
        return max(0, (corte - self.vto_mas_antiguo).days)

    def a_fila(self, corte: dt.date) -> list[object]:
        return [
            self.nombre,
            self.telefono,
            self.contrato,
            self.cuotas,
            round(self.total_vencido, 2),
            round(self.pagado_a_cuenta, 2),
            ", ".join(self.meses),
            self.vto_mas_antiguo,
            self.dias_de_mora(corte),
            self.localidad,
            self.provincia,
            self.supervisor,
            self.vendedor,
        ]


def _vencimiento(fila: FilaCobranza, anio: int, usar_fin_de_mes: bool) -> dt.date | None:
    if fila.vto_cuota:
        return fila.vto_cuota
    if usar_fin_de_mes:
        return fin_de_mes(fila.hoja, anio)
    return None


def calcular_mora(
    base: dict[str, list[FilaCobranza]],
    cfg: Config,
    anio: int,
    corte: dt.date | None = None,
) -> list[ClienteEnMora]:
    """Clientes en mora al dia de corte, ordenados por monto vencido."""
    corte = corte or dt.date.today()
    gracia = int(cfg.mora.get("dias_gracia", 0))
    incluir_parciales = bool(cfg.mora.get("incluir_parciales", True))
    usar_fin_de_mes = bool(cfg.mora.get("usar_fin_de_mes_si_falta_vto", True))
    parcial = cfg.mora.get("estado_parcial", PARCIAL)
    limite = corte - dt.timedelta(days=gracia)

    por_cliente: dict[str, ClienteEnMora] = {}
    for mes in cfg.meses:
        for fila in base.get(mes, []):
            if fila.es_baja:
                continue
            estado = estado_de(fila.monto_pago, fila.cuota, parcial)
            if estado != PENDIENTE and not (incluir_parciales and estado == parcial):
                continue
            vence = _vencimiento(fila, anio, usar_fin_de_mes)
            if vence is None or vence > limite:
                continue  # todavia no vencio: no es mora
            if fila.saldo <= 0:
                continue

            clave = fila.contrato or f"nombre:{fila.nombre.lower()}"
            cliente = por_cliente.get(clave)
            if cliente is None:
                cliente = ClienteEnMora(
                    contrato=fila.contrato,
                    nombre=fila.nombre,
                    telefono=fila.telefono,
                    localidad=fila.localidad,
                    provincia=fila.provincia,
                    supervisor=fila.supervisor,
                    vendedor=fila.vendedor,
                )
                por_cliente[clave] = cliente
            # Los datos de contacto se completan con la primera fila que los
            # tenga: en el archivo real no todas las pestañas los repiten.
            cliente.telefono = cliente.telefono or fila.telefono
            cliente.localidad = cliente.localidad or fila.localidad
            cliente.provincia = cliente.provincia or fila.provincia
            cliente.supervisor = cliente.supervisor or fila.supervisor
            cliente.vendedor = cliente.vendedor or fila.vendedor

            cliente.cuotas += 1
            cliente.total_vencido += fila.saldo
            cliente.pagado_a_cuenta += fila.monto_pago or 0.0
            if fila.hoja not in cliente.meses:
                cliente.meses.append(fila.hoja)
            if cliente.vto_mas_antiguo is None or vence < cliente.vto_mas_antiguo:
                cliente.vto_mas_antiguo = vence

    return sorted(
        por_cliente.values(), key=lambda c: (-c.total_vencido, c.nombre)
    )


_ROJO = "FFC00000"
_AZUL = "FF1F4E78"
_BLANCO = "FFFFFFFF"


def escribir_tablero(
    clientes: list[ClienteEnMora],
    ruta: str | Path,
    corte: dt.date | None = None,
) -> Path:
    """Escribe el tablero de mora en un .xlsx nuevo, listo para trabajar."""
    corte = corte or dt.date.today()
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.title = "Mora Activa"

    ws.append([f"Mora activa al {corte.strftime('%d/%m/%Y')}"])
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=_ROJO)
    ws.append(
        [
            f"{len(clientes)} clientes",
            None,
            None,
            sum(c.cuotas for c in clientes),
            round(sum(c.total_vencido for c in clientes), 2),
        ]
    )
    for celda in ws[2]:
        celda.font = Font(name="Calibri", size=11, bold=True)
    ws.append([])

    fila_enc = ws.max_row + 1
    ws.append(ENCABEZADOS_MORA)
    relleno = PatternFill("solid", fgColor=_AZUL)
    for celda in ws[fila_enc]:
        celda.font = Font(name="Calibri", size=11, bold=True, color=_BLANCO)
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for cliente in clientes:
        ws.append(cliente.a_fila(corte))

    ws.freeze_panes = f"A{fila_enc + 1}"
    ws.auto_filter.ref = (
        f"A{fila_enc}:{get_column_letter(len(ENCABEZADOS_MORA))}{ws.max_row}"
    )
    for fila in ws.iter_rows(min_row=fila_enc + 1, max_row=ws.max_row):
        fila[4].number_format = "#,##0.00"
        fila[5].number_format = "#,##0.00"
        fila[7].number_format = "dd/mm/yyyy"
    anchos = {"A": 34, "B": 14, "C": 10, "D": 10, "E": 18, "F": 15,
              "G": 30, "H": 14, "I": 11, "J": 22, "K": 14, "L": 20, "M": 20}
    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    ruta = Path(ruta)
    libro.save(ruta)
    libro.close()
    return ruta
