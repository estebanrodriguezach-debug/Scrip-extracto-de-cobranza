"""Lectura y escritura de ``Cobranza General``, la base operativa.

Lo que tiene de particular el archivo real, y que obliga a leerlo con cuidado:

* El encabezado esta en la **fila 2**, no en la 1, y las columnas A, B y C no
  tienen titulo. Por eso las columnas se toman de ``conciliacion.json`` por
  letra, y el encabezado se usa solo para avisar si algo se corrio de lugar.
* Abajo de la tabla de clientes hay un **bloque de estadisticas** en las mismas
  columnas ('Cobrado', 'Pendiente', 'Baja' con totales y porcentajes). Esas
  filas no son clientes y hay que descartarlas: se reconocen porque no tienen
  nombre en la columna B.
* ``Estado`` tiene un tercer valor que no estaba en la especificacion: **3**,
  cargado a mano, que significa baja (620 filas en el año, con la palabra
  'baja' escrita en la columna de fecha de pago). Ni la conciliacion ni el
  tablero de mora lo pisan.
* La columna ``Telefono`` tiene celdas con formato de fecha y un numero
  imposible: openpyxl las reporta como error. Se leen como texto y listo.
"""

from __future__ import annotations

import datetime as dt
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from clasificador.config import Config
from clasificador.reglas import normalizar
from clasificador.xlsx_parche import Celda


def _a_texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, dt.datetime):
        return valor.date().isoformat()
    return str(valor).strip()


def _a_fecha(valor: object) -> dt.date | None:
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    return None


def _a_numero(valor: object) -> float | None:
    if valor is None or valor == "" or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("$", "").replace(" ", "")
    if not texto:
        return None
    # "1.234,56" (como se escribe aca) y "1234.56" tienen que entrar los dos.
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


@dataclass
class FilaCobranza:
    """Una fila de cliente de una hoja mensual."""

    hoja: str
    fila: int
    nombre: str = ""
    contrato: str = ""
    telefono: str = ""
    localidad: str = ""
    provincia: str = ""
    fecha_pago: dt.date | None = None
    texto_fecha_pago: str = ""
    monto_pago: float | None = None
    medio_pago: str = ""
    cuota: float | None = None
    mes_inicio: str = ""
    estado: object = None
    supervisor: str = ""
    vendedor: str = ""
    vto_cuota: dt.date | None = None
    obse: str = ""
    textos_de_baja: tuple[str, ...] = field(default=(), repr=False)
    estado_baja: int = 3

    @property
    def es_baja(self) -> bool:
        """Contrato dado de baja: no se concilia ni se reclama."""
        if _a_numero(self.estado) == float(self.estado_baja):
            return True
        for texto in (self.texto_fecha_pago, self.obse):
            if normalizar(texto) in self.textos_de_baja:
                return True
        return False

    @property
    def tiene_cobro(self) -> bool:
        return self.fecha_pago is not None or bool(self.monto_pago)

    @property
    def pendiente(self) -> bool:
        """Lista para recibir un pago: sin cobro cargado y no dada de baja."""
        return not self.es_baja and not self.tiene_cobro

    @property
    def saldo(self) -> float:
        """Cuanto falta para cubrir la cuota. 0 si esta cubierta o no hay cuota."""
        if self.cuota is None:
            return 0.0
        return max(0.0, self.cuota - (self.monto_pago or 0.0))


def _leer_hoja(ws, cfg: Config) -> list[FilaCobranza]:
    idx = cfg.indices
    primera = cfg.cobranza["primera_fila_datos"]
    maximo = cfg.cobranza["contrato_maximo"]
    textos_baja = tuple(normalizar(t) for t in cfg.cobranza["textos_de_baja"])
    estado_baja = int(cfg.cobranza["estado_baja"])

    def celda(fila: tuple, campo: str) -> object:
        i = idx[campo] - 1
        return fila[i] if i < len(fila) else None

    filas: list[FilaCobranza] = []
    for nro, fila in enumerate(
        ws.iter_rows(min_row=primera, values_only=True), start=primera
    ):
        nombre = _a_texto(celda(fila, "nombre"))
        if not nombre:
            # Sin nombre no es un cliente: es el bloque de estadisticas que
            # vive abajo de la tabla, o una fila en blanco.
            continue
        contrato = _a_texto(celda(fila, "contrato"))
        if contrato:
            numero = _a_numero(contrato)
            contrato = (
                str(int(numero))
                if numero is not None and numero == int(numero) and 0 < numero <= maximo
                else ""
            )
        crudo_fecha = celda(fila, "fecha_pago")
        filas.append(
            FilaCobranza(
                hoja=ws.title,
                fila=nro,
                nombre=nombre,
                contrato=contrato,
                telefono=_a_texto(celda(fila, "telefono")),
                localidad=_a_texto(celda(fila, "localidad")),
                provincia=_a_texto(celda(fila, "provincia")),
                fecha_pago=_a_fecha(crudo_fecha),
                texto_fecha_pago="" if _a_fecha(crudo_fecha) else _a_texto(crudo_fecha),
                monto_pago=_a_numero(celda(fila, "monto_pago")),
                medio_pago=_a_texto(celda(fila, "medio_pago")),
                cuota=_a_numero(celda(fila, "cuota")),
                mes_inicio=_a_texto(celda(fila, "mes_inicio")),
                estado=celda(fila, "estado"),
                supervisor=_a_texto(celda(fila, "supervisor")),
                vendedor=_a_texto(celda(fila, "vendedor")),
                vto_cuota=_a_fecha(celda(fila, "vto_cuota")),
                obse=_a_texto(celda(fila, "obse")),
                textos_de_baja=textos_baja,
                estado_baja=estado_baja,
            )
        )
    return filas


def verificar_encabezado(ws, cfg: Config) -> list[str]:
    """Avisa si una columna no esta donde dice ``conciliacion.json``.

    No corta el proceso: devuelve avisos. La idea es que si alguien inserta una
    columna en el Excel, el operador se entere por un mensaje claro en vez de
    que la conciliacion escriba el monto en la columna de al lado.
    """
    fila_enc = cfg.cobranza["fila_encabezado"]
    leidos = {
        celda.column: normalizar(celda.value)
        for celda in next(ws.iter_rows(min_row=fila_enc, max_row=fila_enc))
        if celda.value
    }
    avisos = []
    for campo, i in cfg.indices.items():
        esperado = normalizar(campo.replace("_", " "))
        real = leidos.get(i, "")
        if not real:
            continue  # A, B y C no tienen titulo en el archivo real
        if esperado.split()[0] not in real and real.split()[0] not in esperado:
            avisos.append(
                f"Hoja {ws.title!r}: la columna de {campo} deberia ser "
                f"{cfg.cobranza['columnas'][campo]} pero ahi dice {real!r}. "
                "Revisa conciliacion.json antes de seguir."
            )
    return avisos


def leer_cobranza(
    ruta: str | Path, cfg: Config, meses: list[str] | None = None
) -> tuple[dict[str, list[FilaCobranza]], list[str]]:
    """Lee las hojas mensuales. Devuelve ({mes: filas}, avisos).

    Se lee con ``data_only=True`` para tomar el valor calculado de ``Estado`` y
    ``Cantidad cuotas``, que son formulas. La escritura va por otro camino
    (``xlsx_parche``) justamente para no perder esas formulas.
    """
    pedidos = meses or cfg.meses
    libro = openpyxl.load_workbook(ruta, data_only=True)
    try:
        resultado: dict[str, list[FilaCobranza]] = {}
        avisos: list[str] = []
        for mes in pedidos:
            if mes not in libro.sheetnames:
                if meses:  # si lo pidieron explicitamente, hay que avisar
                    avisos.append(f"El libro no tiene la hoja {mes!r}.")
                continue
            ws = libro[mes]
            avisos.extend(verificar_encabezado(ws, cfg))
            resultado[mes] = _leer_hoja(ws, cfg)
        if not resultado:
            raise RuntimeError(
                f"No encontre ninguna hoja mensual en {Path(ruta).name}. "
                f"Buscaba: {', '.join(pedidos)}. "
                f"El libro tiene: {', '.join(libro.sheetnames[:12])}..."
            )
        return resultado, avisos
    finally:
        libro.close()


def anio_de_base(base: dict[str, list[FilaCobranza]]) -> int:
    """Año al que corresponde el libro, deducido de sus propias fechas.

    ``Cobranza General 2026`` solo tiene las cuotas de 2026, pero el nombre del
    archivo no es confiable como fuente de verdad. Se mira el año que mas
    repiten las fechas de pago y de vencimiento ya cargadas.
    """
    anios: Counter[int] = Counter()
    for filas in base.values():
        for f in filas:
            for fecha in (f.fecha_pago, f.vto_cuota):
                if fecha:
                    anios[fecha.year] += 1
    return anios.most_common(1)[0][0] if anios else 0


def celdas_de_cobro(
    fila: FilaCobranza,
    cfg: Config,
    fecha: dt.date,
    monto: float,
    medio: str,
) -> list[Celda]:
    """Las tres celdas que completa la conciliacion: fecha, monto y medio."""
    return [
        Celda(fila.hoja, fila.fila, cfg.indice("fecha_pago"), fecha),
        Celda(fila.hoja, fila.fila, cfg.indice("monto_pago"), monto),
        Celda(fila.hoja, fila.fila, cfg.indice("medio_pago"), medio),
    ]
