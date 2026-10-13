"""Armado del .xlsx del extracto, con el mismo formato que la planilla de Google.

Para Google Sheets el formato se arma con la API (ver ``sheets.preparar``). Para
SharePoint no hay API equivalente: la Excel REST API de Graph no admite permisos
de aplicacion y encima no expone validacion de datos, asi que no podria crear el
desplegable de Categoria. La salida, entonces, es armar el libro entero aca con
openpyxl y subirlo como archivo; Graph lo guarda tal cual, con hojas,
desplegable y formatos intactos.

Este modulo es el espejo local de ``PlanillaCobranzas.preparar``: si cambia el
formato de una, tiene que cambiar el de la otra.
"""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from clasificador.excel import (
    ENCABEZADOS,
    HOJA_DATOS,
    Movimiento,
    claves_con_ocurrencia,
    leer_movimientos,
)
from clasificador.reglas import JuegoDeReglas

HOJA_LISTAS = "Listas"

# El desplegable se aplica a un rango fijo, no solo a las filas escritas: si no,
# las filas que alguien agregue a mano quedarian sin validacion.
FILAS_MINIMAS_DESPLEGABLE = 5000

# Tomados del Extractos_cobranzas.xlsx que se venia usando, para que el
# archivo generado se vea igual al de antes y no apenas parecido. La B no
# estaba fijada en el original, asi que se deja en el ancho por defecto.
ANCHOS = {"A": 12.0, "C": 15.0, "D": 60.0, "E": 30.0, "F": 38.0}

# Encabezado del original: Calibri 11 en negrita, blanco sobre azul oscuro.
COLOR_ENCABEZADO = "FF1F4E78"
COLOR_TEXTO_ENCABEZADO = "FFFFFFFF"

# El original guarda las fechas en mm-dd-yy. Es ambiguo de leer, pero cambiarlo
# haria que el archivo se vea distinto al que el equipo viene mirando.
FORMATO_FECHA = "mm-dd-yy"
FORMATO_IMPORTE = "#,##0.00"


def _fila(mov: Movimiento) -> list[object]:
    """Como ``Movimiento.a_fila`` pero con fecha e importe nativos.

    ``a_fila`` serializa todo a texto porque Google Sheets lo interpreta al
    escribirlo (``USER_ENTERED``). Excel no: una fecha escrita como texto queda
    como texto, no se ordena ni se filtra por mes. Como el archivo se arma aca,
    se escriben los tipos de verdad; ``leer_movimientos`` los lee igual.
    """
    return [
        mov.fecha,
        mov.banco,
        mov.importe,
        mov.extracto,
        mov.detalle,
        mov.categoria,
    ]


def armar_libro(
    movimientos: list[Movimiento], categorias: tuple[str, ...] | list[str]
) -> bytes:
    """Devuelve los bytes de un .xlsx con el formato del extracto.

    Hoja ``Cobranzas`` con las 6 columnas, hoja ``Listas`` con las categorias,
    desplegable sobre la columna F (Categoria) y fila 1 congelada.
    """
    categorias = list(categorias)
    libro = openpyxl.Workbook()

    ws = libro.active
    ws.title = HOJA_DATOS
    ws.append(ENCABEZADOS)
    for mov in movimientos:
        ws.append(_fila(mov))

    listas = libro.create_sheet(HOJA_LISTAS)
    for categoria in categorias:
        listas.append([categoria])

    validacion = DataValidation(
        type="list",
        # Sin el '=' adelante: en OOXML el contenido de <formula1> es la formula
        # pelada y openpyxl la escribe verbatim. Con el '=' Excel de escritorio
        # y Excel Online ofrecen reparar el archivo y descartan la validacion,
        # que es justamente el desplegable que hay que conservar.
        formula1=f"{HOJA_LISTAS}!$A$1:$A${len(categorias)}",
        allow_blank=True,
        # Trampa de OOXML: el atributo se llama showDropDown pero esta invertido,
        # en True ESCONDE la flechita. False es lo que muestra el desplegable.
        showDropDown=False,
    )
    validacion.errorTitle = "Categoría"
    validacion.error = "Elegí una categoría de la lista."
    ws.add_data_validation(validacion)
    ultima = max(ws.max_row, FILAS_MINIMAS_DESPLEGABLE)
    validacion.add(f"F2:F{ultima}")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    relleno = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
    centrado = Alignment(horizontal="center", vertical="center")
    for celda in ws[1]:
        celda.font = Font(
            name="Calibri", size=11, bold=True, color=COLOR_TEXTO_ENCABEZADO
        )
        celda.fill = relleno
        celda.alignment = centrado
    for columna, ancho in ANCHOS.items():
        ws.column_dimensions[columna].width = ancho
    for fila in ws.iter_rows(min_row=2, max_row=ws.max_row):
        fila[0].number_format = FORMATO_FECHA
        fila[2].number_format = FORMATO_IMPORTE

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def combinar(
    previos: list[Movimiento],
    nuevos: list[Movimiento],
    juego: JuegoDeReglas,
    respetar_categoria_existente: bool = True,
) -> tuple[list[Movimiento], int]:
    """Mezcla lo que ya esta en SharePoint con lo nuevo. Devuelve (todos, agregados).

    Reemplaza al append incremental de Sheets: como el archivo se reescribe
    entero, sin este merge cada corrida borraria las categorias corregidas a
    mano. Funciona porque ``Movimiento.clave`` no depende de detalle ni de
    categoria, justamente para que editarlas no duplique la fila.
    """
    presentes = set(claves_con_ocurrencia(previos))
    agregados = []
    for mov, clave in zip(nuevos, claves_con_ocurrencia(nuevos)):
        if clave in presentes:
            continue
        if not (respetar_categoria_existente and mov.categoria):
            mov.categoria = juego.clasificar(mov.detalle).categoria
        agregados.append(mov)

    # Los nuevos van al final, sin reordenar lo que ya estaba: es lo mismo que
    # hace el append de Google Sheets, asi el archivo no se revuelve entero
    # cada vez que se agregan tres filas.
    return previos + agregados, len(agregados)


def leer_libro_existente(crudo: bytes, nombre: str) -> list[Movimiento]:
    """Lee el .xlsx que ya estaba en el destino, con un error entendible.

    En una carpeta compartida de Administracion y Finanzas es muy posible que
    ya haya un archivo con ese nombre que sea otra cosa: una copia de la base
    de origen (que se llama igual), un export con la pestana 'Hoja1', o algo
    que ni siquiera es un .xlsx. Sin esto el usuario recibe un KeyError o un
    BadZipFile crudos, justo despues de que el script le bajo el archivo.
    """
    try:
        return leer_movimientos(io.BytesIO(crudo))
    except Exception as error:
        raise RuntimeError(
            f"El archivo {nombre!r} que ya esta en la carpeta no tiene el "
            f"formato del extracto: {_describir_libro(crudo)}.\n"
            "Renombralo en SharePoint, o corre el comando con "
            "--nombre OTRO-NOMBRE.xlsx para no pisarlo."
        ) from error


def _describir_libro(crudo: bytes) -> str:
    """Dice por que no se pudo leer: que hojas tiene, o que ni es un .xlsx."""
    try:
        libro = openpyxl.load_workbook(io.BytesIO(crudo), read_only=True)
    except Exception:
        return "no es un .xlsx valido"
    try:
        hojas = ", ".join(repr(h) for h in libro.sheetnames)
    finally:
        libro.close()
    return f"no tiene la hoja {HOJA_DATOS!r}, tiene {hojas}"


def volcar_en_carpeta(
    carpeta: str | Path,
    movimientos: list[Movimiento],
    juego: JuegoDeReglas,
    nombre: str = "Extractos_cobranzas.xlsx",
    respetar_categoria_existente: bool = True,
) -> tuple[Path, int, int]:
    """Escribe el extracto en una carpeta del disco. Devuelve (ruta, agregados, previos).

    Pensado para una carpeta de SharePoint sincronizada con OneDrive: el
    archivo se guarda local y OneDrive lo sube solo. Evita todo el tramite de
    Entra ID, que necesita que un administrador del tenant autorice la app.

    Si el archivo ya esta, se lo lee primero y se mezclan los movimientos
    nuevos, para no pisar las categorias corregidas a mano.
    """
    carpeta = Path(carpeta).expanduser()
    if not carpeta.is_dir():
        raise RuntimeError(
            f"La carpeta {carpeta} no existe. Si es la de OneDrive, revisa que "
            "la biblioteca de SharePoint este sincronizada y que la ruta sea la "
            "que muestra el Explorador."
        )

    destino = carpeta / nombre
    previos: list[Movimiento] = []
    if destino.exists():
        previos = leer_libro_existente(destino.read_bytes(), nombre)

    todos, agregados = combinar(
        previos, movimientos, juego, respetar_categoria_existente
    )

    # Se escribe a un temporal en la misma carpeta y recien despues se renombra:
    # si no, OneDrive puede empezar a subir un archivo a medio escribir.
    temporal = carpeta / f".{nombre}.tmp"
    temporal.write_bytes(armar_libro(todos, juego.categorias_validas))
    temporal.replace(destino)
    return destino, agregados, len(previos)
